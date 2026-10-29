#!/usr/bin/env python3
"""
reportprep: turn scan output into a report-prep spreadsheet.

Currently reads .nessus files. Maps plugins to your library titles via
mapping.toml, consolidates affected IPs into ranges, picks a
representative host for evidence, and suggests a verification command.

Usage:
    reportprep.py --eng-root "$ENG_ROOT"          # reads $ENG_ROOT/nessus/*.nessus
    reportprep.py scan1.nessus scan2.nessus -o prep.xlsx

Everything runs locally. Nothing is sent anywhere.
"""
from __future__ import annotations

import argparse
import ipaddress
import re
import sys
import tomllib
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is missing: pip install openpyxl  (or apt install python3-openpyxl)")

HERE = Path(__file__).resolve().parent
SEV_NAMES = {0: "Info", 1: "Low", 2: "Medium", 3: "High", 4: "Critical"}
SEV_FILL = {4: "7B1F2A", 3: "C0392B", 2: "E67E22", 1: "2E86C1", 0: "7F8C8D"}
AREAS = [
    "Secure Configuration",
    "Vulnerability Management",
    "Encryption",
    "Identity & Authentication",
    "Access Control",
]
HTTPS_PORTS = {443, 4443, 8443, 9443}
EVIDENCE_MAX = 1500


# --------------------------------------------------------------------- data

@dataclass
class Item:
    host: str
    port: int
    proto: str
    svc: str
    plugin_id: int
    plugin_name: str
    severity: int
    output: str
    source: str


@dataclass
class Rule:
    title: str
    area: str
    ids: set[int]
    patterns: list[re.Pattern]
    verify: str
    include_info: bool


@dataclass
class Group:
    title: str
    area: str
    verify: str
    mapped: bool
    items: list[Item] = field(default_factory=list)

    @property
    def severity(self) -> int:
        return max(i.severity for i in self.items)

    @property
    def hosts(self) -> list[str]:
        return sorted({i.host for i in self.items}, key=host_sort_key)


# ------------------------------------------------------------------ mapping

def load_rules(path: Path) -> list[Rule]:
    with open(path, "rb") as fh:
        data = tomllib.load(fh)
    rules = []
    for f in data.get("finding", []):
        rules.append(Rule(
            title=f["title"],
            area=f.get("area", ""),
            ids=set(f.get("nessus_ids", [])),
            patterns=[re.compile(p, re.I) for p in f.get("nessus_name_patterns", [])],
            verify=f.get("verify", ""),
            include_info=f.get("include_info", False),
        ))
    return rules


def match(item: Item, rules: list[Rule]) -> Rule | None:
    for r in rules:
        if item.plugin_id in r.ids and (item.severity > 0 or r.include_info):
            return r
    for r in rules:
        if any(p.search(item.plugin_name) for p in r.patterns):
            if item.severity > 0 or r.include_info:
                return r
    return None


# ------------------------------------------------------------------ parsing

def parse_nessus(path: Path) -> list[Item]:
    items = []
    for _, host in ET.iterparse(path, events=("end",)):
        if host.tag != "ReportHost":
            continue
        name = host.get("name", "")
        for ri in host.findall("ReportItem"):
            items.append(Item(
                host=name,
                port=int(ri.get("port", "0") or 0),
                proto=ri.get("protocol", ""),
                svc=ri.get("svc_name", ""),
                plugin_id=int(ri.get("pluginID", "0") or 0),
                plugin_name=ri.get("pluginName", ""),
                severity=int(ri.get("severity", "0") or 0),
                output=(ri.findtext("plugin_output") or "").strip(),
                source=path.name,
            ))
        host.clear()
    return items


def dedupe(items: list[Item]) -> list[Item]:
    seen, out = set(), []
    for i in items:
        key = (i.host, i.port, i.proto, i.plugin_id)
        if key not in seen:
            seen.add(key)
            out.append(i)
    return out


# ------------------------------------------------------------ IP handling

def host_sort_key(h: str):
    try:
        return (0, int(ipaddress.ip_address(h)), h)
    except ValueError:
        return (1, 0, h.lower())


def consolidate(hosts) -> str:
    """'1.1.1.1-10, 1.1.1.15, web.example' style ranges."""
    v4, other = [], []
    for h in set(hosts):
        try:
            ip = ipaddress.ip_address(h)
            (v4 if ip.version == 4 else other).append(ip if ip.version == 4 else h)
        except ValueError:
            other.append(h)
    v4.sort()
    parts, i = [], 0
    while i < len(v4):
        start = end = v4[i]
        while (i + 1 < len(v4)
               and int(v4[i + 1]) == int(end) + 1
               and str(v4[i + 1]).rsplit(".", 1)[0] == str(start).rsplit(".", 1)[0]):
            i += 1
            end = v4[i]
        if start == end:
            parts.append(str(start))
        else:
            parts.append(f"{start}-{str(end).rsplit('.', 1)[1]}")
        i += 1
    parts += sorted((str(o) for o in other), key=host_sort_key)
    return ", ".join(parts)


# ------------------------------------------------------------- formatting

def url_for(item: Item) -> str:
    # Prefer a full URL Nessus already printed in the plugin output
    m = re.search(r"https?://[^\s'\"<>]+", item.output)
    if m:
        return m.group(0)
    scheme = "https" if item.port in HTTPS_PORTS or "https" in item.svc else "http"
    return f"{scheme}://{item.host}:{item.port}/"


def verify_cmd(template: str, item: Item) -> str:
    if not template:
        return f"nmap -sV -p{item.port} {item.host}" if item.port else ""
    return template.format(host=item.host, port=item.port, url=url_for(item))


def trim(text: str) -> str:
    return text if len(text) <= EVIDENCE_MAX else text[:EVIDENCE_MAX] + "\n--- <snip> ---"


def suggested_rating(worst: int | None) -> str:
    if worst is None:
        return "Good"
    return {0: "Good", 1: "Good", 2: "Fair", 3: "Poor", 4: "Poor"}[worst]


# ---------------------------------------------------------------- workbook

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    return ws


def sev_cell(ws, row, col, sev):
    cell = ws.cell(row=row, column=col, value=SEV_NAMES[sev])
    cell.fill = PatternFill("solid", fgColor=SEV_FILL[sev])
    cell.font = Font(bold=True, color="FFFFFF")


def finish(ws):
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def read_kv(path: Path) -> dict:
    out = {}
    if path.exists():
        for line in path.read_text().splitlines():
            if "=" in line:
                k, v = line.split("=", 1)
                out[k.strip()] = v.strip()
    return out


def build(groups, unmapped, out_path: Path, eng_root: Path | None):
    wb = Workbook()
    wb.remove(wb.active)
    ordered = sorted(groups, key=lambda g: (-g.severity, -len(g.hosts), g.title))

    # ---- Summary
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60
    bold = Font(bold=True)
    info = read_kv(eng_root / "engagement.info") if eng_root else {}
    r = 1
    if info:
        ws.cell(row=r, column=1, value="2.4 Information Provided").font = bold
        r += 1
        line = "Low-level domain credentials provided"
        if info.get("username"):
            dom = info.get("domain", "")
            line += f", username: {dom + chr(92) if dom else ''}{info['username']}"
        ws.cell(row=r, column=1, value=line)
        r += 2

    ws.cell(row=r, column=1, value="Findings by risk").font = bold
    r += 1
    for sev in (4, 3, 2, 1, 0):
        n = sum(1 for g in groups if g.severity == sev)
        if n:
            sev_cell(ws, r, 1, sev)
            ws.cell(row=r, column=2, value=n)
            r += 1
    r += 1

    ws.cell(row=r, column=1, value="1.4 Overall Assessment (suggested)").font = bold
    r += 1
    for area in AREAS:
        sevs = [g.severity for g in groups if g.area == area]
        worst = max(sevs) if sevs else None
        ws.cell(row=r, column=1, value=area)
        note = f"{len(sevs)} finding(s), worst {SEV_NAMES[worst]}" if sevs else "no scanner findings"
        ws.cell(row=r, column=2, value=f"{suggested_rating(worst)}  ({note})")
        r += 1
    ws.cell(row=r, column=1, value="Suggestions only: worst risk High/Critical → Poor, Medium → Fair, else Good.").font = Font(italic=True, color="7F8C8D")
    r += 2

    ws.cell(row=r, column=1, value="3.3 Priority order").font = bold
    r += 1
    for n, g in enumerate([g for g in ordered if g.severity > 0][:5], 1):
        ws.cell(row=r, column=1, value=f"{n}. {g.title}")
        ws.cell(row=r, column=2, value=f"{SEV_NAMES[g.severity]}, {len(g.hosts)} host(s)")
        r += 1

    # ---- Findings
    ws = sheet(wb, "Findings",
               ["Title", "In library?", "Risk", "Area", "Affected IPs", "Host count",
                "Nessus plugins", "Verification command (representative host)",
                "Evidence (representative host)", "Notes"],
               [38, 11, 10, 22, 34, 10, 40, 48, 70, 30])
    for row, g in enumerate(ordered, 2):
        rep = max(g.items, key=lambda i: (i.severity, len(i.output)))
        plugins = sorted({f"{i.plugin_id}: {i.plugin_name}" for i in g.items})
        ws.cell(row=row, column=1, value=g.title)
        ws.cell(row=row, column=2, value="Yes" if g.mapped else "No")
        sev_cell(ws, row, 3, g.severity)
        ws.cell(row=row, column=4, value=g.area)
        ws.cell(row=row, column=5, value=consolidate(g.hosts))
        ws.cell(row=row, column=6, value=len(g.hosts))
        ws.cell(row=row, column=7, value="\n".join(plugins))
        ws.cell(row=row, column=8, value=verify_cmd(g.verify, rep))
        ws.cell(row=row, column=9, value=f"[{rep.host}:{rep.port}]\n{trim(rep.output)}" if rep.output else f"[{rep.host}:{rep.port}]")
    finish(ws)

    # ---- Affected (one row per finding/host/port)
    ws = sheet(wb, "Affected",
               ["Title", "Host", "Port", "Proto", "Service", "Risk", "Plugin ID",
                "Plugin name", "Verification command"],
               [38, 16, 7, 7, 12, 10, 10, 50, 55])
    row = 2
    for g in ordered:
        for i in sorted(g.items, key=lambda i: (host_sort_key(i.host), i.port, i.plugin_id)):
            ws.cell(row=row, column=1, value=g.title)
            ws.cell(row=row, column=2, value=i.host)
            ws.cell(row=row, column=3, value=i.port)
            ws.cell(row=row, column=4, value=i.proto)
            ws.cell(row=row, column=5, value=i.svc)
            sev_cell(ws, row, 6, i.severity)
            ws.cell(row=row, column=7, value=i.plugin_id)
            ws.cell(row=row, column=8, value=i.plugin_name)
            ws.cell(row=row, column=9, value=verify_cmd(g.verify, i))
            row += 1
    finish(ws)

    # ---- Unmapped
    ws = sheet(wb, "Unmapped",
               ["Plugin ID", "Nessus name", "Risk", "Affected IPs", "Host count",
                "Verification command (representative host)", "Tenable link", "Map to library title"],
               [10, 50, 10, 34, 10, 40, 45, 34])
    by_plugin = defaultdict(list)
    for i in unmapped:
        by_plugin[(i.plugin_id, i.plugin_name)].append(i)
    rows = sorted(by_plugin.items(), key=lambda kv: (-max(i.severity for i in kv[1]), kv[0][1]))
    for row, ((pid, pname), its) in enumerate(rows, 2):
        rep = its[0]
        ws.cell(row=row, column=1, value=pid)
        ws.cell(row=row, column=2, value=pname)
        sev_cell(ws, row, 3, max(i.severity for i in its))
        ws.cell(row=row, column=4, value=consolidate(i.host for i in its))
        ws.cell(row=row, column=5, value=len({i.host for i in its}))
        ws.cell(row=row, column=6, value=verify_cmd("", rep))
        ws.cell(row=row, column=7, value=f"https://www.tenable.com/plugins/nessus/{pid}")
    finish(ws)

    # ---- Scope
    if eng_root:
        targets = []
        for name in ("scope.txt", "targets.txt"):
            p = eng_root / name
            if p.exists():
                targets += [l.strip() for l in p.read_text().splitlines()
                            if l.strip() and not l.strip().startswith("#")]
        if targets:
            ws = sheet(wb, "Scope", ["2.1 Scope of Testing (consolidated)", "Raw entries"], [50, 30])
            singles = consolidate(t for t in targets if "/" not in t and "-" not in t)
            ranges = list(dict.fromkeys(t for t in targets if "/" in t or "-" in t))
            ws.cell(row=2, column=1, value=", ".join(x for x in [singles, *ranges] if x))
            for r, t in enumerate(dict.fromkeys(targets), 2):
                ws.cell(row=r, column=2, value=t)
            finish(ws)

    # ---- Limitations notes
    if eng_root and (eng_root / "notes" / "limitations.md").exists():
        notes = (eng_root / "notes" / "limitations.md").read_text().strip()
        if notes:
            ws = sheet(wb, "Limitations", ["2.5 Limitations & Constraints (from notes)"], [100])
            for r, line in enumerate(notes.splitlines(), 2):
                ws.cell(row=r, column=1, value=line)

    wb.save(out_path)


# --------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="Build a report-prep spreadsheet from .nessus files.")
    ap.add_argument("files", nargs="*", type=Path, help=".nessus files")
    ap.add_argument("--eng-root", type=Path, help="engagement folder (reads nessus/*.nessus, scope, engagement.info)")
    ap.add_argument("--mapping", type=Path, default=HERE / "mapping.toml")
    ap.add_argument("-o", "--out", type=Path)
    a = ap.parse_args()

    files = list(a.files)
    if a.eng_root:
        files += sorted((a.eng_root / "nessus").glob("*.nessus"))
    if not files:
        ap.error("no .nessus files given or found in --eng-root/nessus/")

    rules = load_rules(a.mapping)
    items = dedupe([i for f in files for i in parse_nessus(f)])

    groups: dict[str, Group] = {}
    unmapped = []
    for i in items:
        r = match(i, rules)
        if r:
            g = groups.setdefault(r.title, Group(r.title, r.area, r.verify, True))
            g.items.append(i)
        elif i.severity > 0:
            unmapped.append(i)

    out = a.out or ((a.eng_root or Path(".")) / f"report_prep_{date.today()}.xlsx")
    build(list(groups.values()), unmapped, out, a.eng_root)

    n_info = sum(1 for i in items if i.severity == 0)
    print(f"Read {len(files)} file(s), {len(items)} items ({n_info} informational skipped unless mapped)")
    print(f"Mapped findings: {len(groups)}   Unmapped plugins: {len({i.plugin_id for i in unmapped})}")
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
